from io import BytesIO
from tempfile import TemporaryDirectory

from django.core.files.base import ContentFile
from django.test import override_settings
from rest_framework import status
from rest_framework.test import APIClient, APITestCase
from openpyxl import Workbook, load_workbook

from apps.projects.models import Project, ProjectMember
from apps.testcases.models import TestCase as TCModel, TestCaseImportRecord, TestCaseModule, TestCaseStep, TestCaseComment
from apps.testcases.services import TestCaseExcelImportService, TestCaseImportTemplateService
from apps.users.models import User

# 注意：testcase-list / testcase-detail URL name 与 ui_automation 模块的 router 冲突，
# 这里直接使用路径常量以避免 reverse() 解析到错误的路由。
TC_LIST_URL = '/api/testcases/'


def tc_detail_url(pk):
    return f'/api/testcases/{pk}/'


class TestCaseApiTests(APITestCase):
    """Testcases 模块 - 测试用例 API 冒烟测试"""

    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user(username='tc_tester', password='pass123456')
        self.other_user = User.objects.create_user(username='tc_other', password='pass123456')
        self.project = Project.objects.create(name='用例测试项目', owner=self.user)
        ProjectMember.objects.create(project=self.project, user=self.user, role='tester')
        self.client.force_authenticate(self.user)
        self.list_url = TC_LIST_URL

    def test_list_requires_authentication(self):
        """未认证用户无法访问用例列表"""
        anon_client = APIClient()
        response = anon_client.get(self.list_url)
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_create_testcase(self):
        """创建测试用例"""
        payload = {
            'title': '登录功能验证',
            'description': '验证正常登录流程',
            'preconditions': '用户已注册',
            'steps': '1. 打开登录页\n2. 输入账号密码\n3. 点击登录',
            'expected_result': '登录成功进入首页',
            'priority': 'high',
            'test_type': 'functional',
            'project_id': self.project.id,
        }
        response = self.client.post(self.list_url, payload, format='json')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['title'], '登录功能验证')
        self.assertEqual(TCModel.objects.count(), 1)

    def test_list_testcases_with_pagination(self):
        """列表查询支持分页"""
        for i in range(15):
            TCModel.objects.create(
                title=f'用例{i}', expected_result='通过', project=self.project, author=self.user,
            )

        response = self.client.get(self.list_url, {'page': 1, 'page_size': 5})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data['results']), 5)
        self.assertEqual(response.data['count'], 15)

    def test_retrieve_testcase_detail(self):
        """获取用例详情"""
        tc = TCModel.objects.create(
            title='详情测试', expected_result='通过', project=self.project, author=self.user,
        )
        TestCaseStep.objects.create(testcase=tc, step_number=1, action='点击按钮', expected='跳转')
        detail_url = tc_detail_url(tc.id)
        response = self.client.get(detail_url)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['title'], '详情测试')
        self.assertEqual(len(response.data['step_details']), 1)

    def test_update_testcase(self):
        """更新测试用例"""
        tc = TCModel.objects.create(
            title='旧标题', expected_result='通过', project=self.project, author=self.user,
        )
        detail_url = tc_detail_url(tc.id)
        response = self.client.patch(detail_url, {'title': '新标题'}, format='json')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        tc.refresh_from_db()
        self.assertEqual(tc.title, '新标题')

    def test_delete_testcase(self):
        """删除测试用例"""
        tc = TCModel.objects.create(
            title='待删除', expected_result='通过', project=self.project, author=self.user,
        )
        detail_url = tc_detail_url(tc.id)
        response = self.client.delete(detail_url)
        self.assertEqual(response.status_code, status.HTTP_204_NO_CONTENT)
        self.assertEqual(TCModel.objects.count(), 0)

    def test_search_testcase_by_title(self):
        """按标题搜索用例"""
        TCModel.objects.create(
            title='支付流程验证', expected_result='通过', project=self.project, author=self.user,
        )
        TCModel.objects.create(
            title='登录流程验证', expected_result='通过', project=self.project, author=self.user,
        )
        response = self.client.get(self.list_url, {'search': '支付'})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data['results']), 1)
        self.assertIn('支付', response.data['results'][0]['title'])

    def test_filter_by_priority(self):
        """按优先级过滤用例"""
        TCModel.objects.create(
            title='高优用例', expected_result='通过', project=self.project, author=self.user, priority='high',
        )
        TCModel.objects.create(
            title='低优用例', expected_result='通过', project=self.project, author=self.user, priority='low',
        )
        response = self.client.get(self.list_url, {'priority': 'high'})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data['results']), 1)
        self.assertEqual(response.data['results'][0]['priority'], 'high')

    def test_only_project_members_see_testcases(self):
        """非项目成员无法看到用例"""
        TCModel.objects.create(
            title='内部用例', expected_result='通过', project=self.project, author=self.user,
        )
        self.client.force_authenticate(self.other_user)
        response = self.client.get(self.list_url)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data['results']), 0)


class TestCaseModuleApiTests(APITestCase):
    """模块树层级、移动和用例归属的接口验证。"""

    def setUp(self):
        self.user = User.objects.create_user(username='module_tester', password='pass123456')
        self.project = Project.objects.create(name='模块测试项目', owner=self.user)
        self.other_project = Project.objects.create(name='其他项目', owner=self.user)
        self.client.force_authenticate(self.user)

    def create_module(self, project=None, parent=None, name='模块'):
        return TestCaseModule.objects.create(project=project or self.project, parent=parent, name=name)

    def test_module_creation_enforces_eight_level_limit(self):
        parent = None
        for level in range(1, 9):
            response = self.client.post('/api/testcases/modules/', {
                'project_id': self.project.id, 'parent_id': parent, 'name': f'第{level}层',
            }, format='json')
            self.assertEqual(response.status_code, status.HTTP_201_CREATED)
            parent = response.data['id']

        response = self.client.post('/api/testcases/modules/', {
            'project_id': self.project.id, 'parent_id': parent, 'name': '第9层',
        }, format='json')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_module_move_rejects_cycle_and_depth_overflow(self):
        root = self.create_module(name='根')
        child = self.create_module(parent=root, name='子')
        response = self.client.put(f'/api/testcases/modules/{root.id}/', {'parent_id': child.id}, format='json')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

        deep_parent = None
        for level in range(1, 8):
            deep_parent = self.create_module(name=f'深{level}') if deep_parent is None else self.create_module(parent=deep_parent, name=f'深{level}')
        response = self.client.put(f'/api/testcases/modules/{root.id}/', {'parent_id': deep_parent.id}, format='json')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_module_move_rejects_cross_project_parent(self):
        module = self.create_module(name='待移动')
        foreign_parent = self.create_module(project=self.other_project, name='外部模块')
        response = self.client.put(f'/api/testcases/modules/{module.id}/', {'parent_id': foreign_parent.id}, format='json')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_module_filter_includes_descendants_and_cases_can_move_to_unassigned(self):
        parent = self.create_module(name='父模块')
        child = self.create_module(parent=parent, name='子模块')
        case = TCModel.objects.create(title='子模块用例', expected_result='通过', project=self.project, author=self.user, module=child)
        response = self.client.get(TC_LIST_URL, {'project': self.project.id, 'module': parent.id, 'include_children': 'true'})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['count'], 1)

        response = self.client.post('/api/testcases/move/', {'testcase_ids': [case.id], 'module_id': None}, format='json')
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        case.refresh_from_db()
        self.assertIsNone(case.module_id)


class TestCaseImportTemplateTests(APITestCase):
    def test_template_matches_module_columns_and_case_fields(self):
        workbook = load_workbook(TestCaseImportTemplateService.build_template(), read_only=True)
        headers = list(next(workbook['用例导入模板'].iter_rows(min_row=1, max_row=1, values_only=True)))
        self.assertEqual(headers[:3], [f'{level}级模块' for level in ('一', '二', '三')])
        self.assertEqual(headers[3:], ['用例名称', '优先级', '测试类型', '前置条件', '步骤描述', '预期结果', '备注', '维护人'])


class TestCaseModuleImportTests(APITestCase):
    def setUp(self):
        self.user = User.objects.create_user(username='module_importer', password='pass123456')
        self.project = Project.objects.create(name='模块导入项目', owner=self.user)

    def test_import_builds_and_reuses_module_hierarchy_from_eight_level_columns(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = '用例导入模板'
        module_headers = [f'{level}级模块' for level in ('一', '二', '三', '四', '五', '六', '七', '八')]
        worksheet.append(module_headers + TestCaseImportTemplateService.HEADERS[3:])
        full_path = [f'模块{level}' for level in range(1, 9)]
        case_fields = ['高', '功能测试', '', '步骤描述', '预期结果', '备注内容', '']
        worksheet.append(full_path + ['完整路径用例1'] + case_fields)
        worksheet.append(full_path + ['完整路径用例2'] + case_fields)
        worksheet.append(['模块1'] + [''] * 7 + ['一级模块用例'] + ['中', '功能测试', '', '步骤描述', '预期结果', '', ''])
        worksheet.append([''] * 8 + ['无模块用例'] + ['低', '功能测试', '', '步骤描述', '预期结果', '', ''])
        worksheet.append(['跳级父模块', '', '跳级子模块'] + [''] * 5 + ['跳级用例'] + ['中', '功能测试', '', '步骤描述', '预期结果', '', ''])
        output = BytesIO()
        workbook.save(output)

        with TemporaryDirectory() as media_root, override_settings(MEDIA_ROOT=media_root):
            record = TestCaseImportRecord.objects.create(
                import_no='IMP_MODULE_PATH_TEST',
                project=self.project,
                created_by=self.user,
            )
            record.import_file.save('module-paths.xlsx', ContentFile(output.getvalue()), save=True)
            summary = TestCaseExcelImportService.import_record(record)

        self.assertEqual((summary.total_rows, summary.success_count, summary.failed_count), (5, 4, 1))
        self.assertIn('模块层级不能跳级填写', summary.failure_details[0]['errors'])
        self.assertEqual(TestCaseModule.objects.filter(project=self.project).count(), 8)

        leaf = TestCaseModule.objects.get(project=self.project, name='模块8')
        self.assertEqual(leaf.parent.name, '模块7')
        self.assertEqual(leaf.parent.parent.name, '模块6')
        full_path_cases = TCModel.objects.filter(title__startswith='完整路径用例')
        self.assertEqual(full_path_cases.count(), 2)
        self.assertEqual(set(full_path_cases.values_list('module_id', flat=True)), {leaf.id})
        self.assertEqual(set(full_path_cases.values_list('priority', flat=True)), {'high'})
        self.assertEqual(TCModel.objects.get(title='一级模块用例').module.name, '模块1')
        self.assertEqual(TCModel.objects.get(title='一级模块用例').priority, 'medium')
        self.assertIsNone(TCModel.objects.get(title='无模块用例').module_id)
        self.assertEqual(TCModel.objects.get(title='无模块用例').priority, 'low')
